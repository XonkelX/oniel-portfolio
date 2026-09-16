from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    ROOT
    / "docs"
    / "professional-positioning"
    / "Oniel_Alejo_Feliz_Software_Engineer_Resume.docx"
)

FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(62, 69, 78)


def set_font(run, size: float, *, bold: bool = False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def remove_paragraph_border(paragraph_or_style):
    properties = paragraph_or_style._element.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def add_hyperlink(paragraph, text: str, url: str, *, size: float = 9.0):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    properties.extend([fonts, color, underline, size_element])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def paragraph(doc, text: str = "", *, size: float = 9.6, bold: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    if text:
        set_font(p.add_run(text), size, bold=bold)
    return p


def section(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6.5)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title.upper()), 10.0, bold=True)


def role_line(doc, left: str, right: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(0.4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(left), 9.7, bold=True)
    set_font(p.add_run(f" | {right}"), 9.2, color=MUTED)


def bullet(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_after = Pt(1.4)
    p.paragraph_format.line_spacing = 1.06
    set_font(p.add_run(f"- {text}"), 9.3)
    return p


def build():
    doc = Document()
    section_settings = doc.sections[0]
    section_settings.page_width = Inches(8.5)
    section_settings.page_height = Inches(11)
    section_settings.top_margin = Inches(0.58)
    section_settings.bottom_margin = Inches(0.42)
    section_settings.left_margin = Inches(0.55)
    section_settings.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = BLACK
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = BLACK
    remove_paragraph_border(title_style)

    name = doc.add_paragraph()
    name.style = doc.styles["Title"]
    remove_paragraph_border(name)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(0)
    name.paragraph_format.line_spacing = 1.0
    set_font(name.add_run("ONIEL ALEJO FELIZ"), 17, bold=True)

    headline = paragraph(doc)
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(headline.add_run("SOFTWARE ENGINEER"), 10.5, bold=True)
    set_font(
        headline.add_run("  |  TypeScript | Node.js | React / Next.js | PostgreSQL"),
        9.6,
        color=MUTED,
    )

    contact = paragraph(doc)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(contact.add_run("Tampa, FL | 656-200-7791 | Onielbf10@gmail.com"), 9.1)

    links = paragraph(doc)
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(links, "onielalejofeliz.space", "https://onielalejofeliz.space/", size=9.0)
    set_font(links.add_run(" | "), 9.0, color=MUTED)
    add_hyperlink(links, "github.com/XonkelX", "https://github.com/XonkelX", size=9.0)
    set_font(links.add_run(" | "), 9.0, color=MUTED)
    add_hyperlink(
        links,
        "linkedin.com/in/oniel-alejo-feliz-45b293312",
        "https://www.linkedin.com/in/oniel-alejo-feliz-45b293312",
        size=9.0,
    )

    authorization = paragraph(doc)
    authorization.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        authorization.add_run("U.S. permanent resident | No employer sponsorship required | English and Spanish"),
        8.9,
        color=MUTED,
    )

    section(doc, "Summary")
    paragraph(
        doc,
        "Software engineer with merged upstream work in established codebases, focused on TypeScript systems, PostgreSQL data boundaries, reliability, and automated testing. Work across Reticle, Apache Maka, Apache Fineract, Code.Sydney / BlueHex, and Clarvia includes substantive maintainer review and regression coverage.",
        size=9.4,
    )

    section(doc, "Open-Source Engineering")
    role_line(doc, "Open-Source Software Contributor - Independent", "2026-Present")
    bullet(
        doc,
        "Reticle #278: added bounded, batched learned-route persistence across crawls and navigation; preserved project state and exact optional semantics; revised the design after substantive maintainer review and merged after upstream validation.",
    )
    bullet(
        doc,
        "Apache Maka #2989: stopped unbounded MCP tool rediscovery while preserving callable-tool snapshots; revised the design through multiple rounds of edge-case review and added slow-notification and snapshot-preservation regressions before merge.",
    )
    bullet(
        doc,
        "Apache Fineract #475 and #431: migrated the remaining 13 client-area Jasmine/Karma specs (73 tests) to Vitest and added translation-backed titles to 28 Accounting routes; both merged after upstream validation.",
    )
    bullet(
        doc,
        "BlueHex #29 and Clarvia #269: added production-build Playwright coverage and fixed keyboard and screen-reader barriers; incorporated maintainer feedback before both changes merged.",
    )

    section(doc, "Selected Engineering Projects")
    role_line(doc, "Relay - Webhook Delivery and Failure Recovery", "TypeScript | React | Cloudflare Workers | D1 | Queues")
    bullet(
        doc,
        "Implemented a transactional outbox, durable scheduling, signed requests, deterministic retries, lease recovery, encrypted endpoint secrets, replay lineage, and inspectable failure evidence. Source: github.com/XonkelX/relay-webhook-delivery",
    )
    role_line(doc, "Next - Real-Time Queue and Authorization", "Next.js | Supabase | PostgreSQL | Playwright")
    bullet(
        doc,
        "Enforced RLS, private/public data boundaries, transactional commands, idempotency, and concurrency in PostgreSQL; synchronized customer, staff, and public views with Realtime plus authoritative recovery. Source: github.com/XonkelX/next-queue",
    )

    section(doc, "Additional Experience")
    role_line(doc, "AI Training Contributor - Outlier", "Remote | Jul 2024-Feb 2025")
    bullet(
        doc,
        "Completed varied AI-training work under detailed, changing evaluation criteria using independent reasoning, careful guideline interpretation, and consistent quality control.",
    )
    role_line(doc, "Computer Repair Technician - Local Computer Repair Shop", "2024 | Six months")
    bullet(
        doc,
        "Diagnosed operating-system, application, and hardware failures; completed software and component repairs; explained practical solutions clearly to customers.",
    )

    section(doc, "Technical Skills")
    paragraph(
        doc,
        "TypeScript, JavaScript, SQL, React, Next.js, Node.js, REST APIs, PostgreSQL, Supabase, RLS, Prisma, Auth.js, Cloudflare Workers, D1, Queues, Hono, Vitest, Testing Library, Playwright, pgTAP, GitHub Actions, Docker, Vercel, accessibility",
        size=9.2,
    )

    core = doc.core_properties
    core.title = "Oniel Alejo Feliz Software Engineer Resume"
    core.subject = "Software engineering resume focused on TypeScript, Node.js, React, Next.js, and PostgreSQL"
    core.author = "Oniel Alejo Feliz"
    core.keywords = "software engineer, TypeScript, Node.js, React, Next.js, PostgreSQL, open source"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
